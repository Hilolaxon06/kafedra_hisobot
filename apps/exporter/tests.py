"""
Экспорт: секции листов растут под число записей (ничего не обрезается),
свидетельства на программы/БД раскладываются по колонкам 2,4, архив
собирает документы по папкам преподавателей и листов.
"""

import datetime
import io
import shutil
import tempfile
import zipfile
from pathlib import Path

import openpyxl
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings

from apps.core.models import Department, Faculty, University
from apps.dictionaries.models import AcademicDegree, EmploymentType
from apps.reports.models import Funding, Patent, PatentAuthor, Publication
from apps.staff.models import Teacher, TeacherPeriodSnapshot
from apps.workflow.models import ReportingPeriod
from .archive import build_package
from .engine import export_report

TMP = tempfile.mkdtemp(prefix="kafedra_test_export_")


@override_settings(MEDIA_ROOT=f"{TMP}/media", EXPORTS_DIR=Path(TMP) / "exports")
class ExportTestBase(TestCase):
    @classmethod
    def setUpTestData(cls):
        uni = University.objects.create(name="Test universiteti")
        fac = Faculty.objects.create(university=uni, name="Test fakulteti")
        cls.dep = Department.objects.create(faculty=fac, name="Test kafedrasi", short_name="TK")
        cls.period = ReportingPeriod.objects.create(
            department=cls.dep, period_label="2026-mart", period_date=datetime.date(2026, 3, 31),
        )
        cls.phd, _ = AcademicDegree.objects.get_or_create(code="phd", defaults={"name": "PhD"})

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TMP, ignore_errors=True)

    def teacher(self, name, employment=EmploymentType.ASOSIY, degree=None):
        teacher = Teacher.objects.create(department=self.dep, full_name=name)
        TeacherPeriodSnapshot.objects.create(
            teacher=teacher, reporting_period=self.period,
            employment_type=EmploymentType.objects.get(code=employment), academic_degree=degree,
        )
        return teacher

    def patent(self, teacher, patent_type=Patent.PatentType.INVENTION, license=False, **extra):
        patent = Patent.objects.create(
            reporting_period=self.period, title=f"{patent_type} {teacher.full_name}",
            patent_type=patent_type, has_license_agreement=license, **extra,
        )
        PatentAuthor.objects.create(patent=patent, teacher=teacher)
        return patent

    def export(self):
        path, _filled, _skipped = export_report(self.dep, self.period)
        return openpyxl.load_workbook(path)


class SectionGrowthTest(ExportTestBase):
    def test_patents_sheet_keeps_every_author(self):
        # в шаблоне под патенты 3 строки (3..5); четвёртый автор не должен пропасть
        names = [f"Patentchi {i}" for i in range(1, 5)]
        for name in names:
            self.patent(self.teacher(name))

        ws = self.export()["2,4"]
        self.assertEqual([ws.cell(row=r, column=2).value for r in range(3, 7)], names)
        self.assertEqual([ws.cell(row=r, column=1).value for r in range(3, 7)], [1, 2, 3, 4])
        # подвал (A6:D6, подпись «Rektor» A7:D7, примечание A8:D8) уехал на строку ниже
        self.assertTrue(str(ws["A8"].value).startswith("Rektor"))
        self.assertEqual({str(r) for r in ws.merged_cells.ranges}, {"A1:D1", "A7:D7", "A8:D8", "A9:D9"})

    def test_two_internal_part_timers_both_listed(self):
        # под «Ichki o'rindosh» в 1,1 одна строка (19)
        self.teacher("Asosiy Odam", degree=self.phd)
        self.teacher("Ichki Birinchi", EmploymentType.ICHKI, degree=self.phd)
        self.teacher("Ichki Ikkinchi", EmploymentType.ICHKI, degree=self.phd)

        wb = self.export()
        ws = wb["1,1"]
        self.assertEqual(ws["B4"].value, "Asosiy Odam")
        self.assertEqual(ws["A18"].value, "Ichki o'rindosh")
        self.assertEqual((ws["B19"].value, ws["B20"].value), ("Ichki Birinchi", "Ichki Ikkinchi"))
        self.assertEqual(ws["B21"].value, "Jami:")

        ws = wb["5"]
        self.assertEqual((ws["B22"].value, ws["B23"].value), ("Ichki Birinchi", "Ichki Ikkinchi"))
        self.assertTrue(str(ws["B24"].value).startswith("Tashqi"))

    def test_total_formula_follows_inserted_rows(self):
        # 2,3: основной штат 3..16 (14 строк), итог C19 = SUM(C3:C18)
        for i in range(1, 16):
            self.teacher(f"Oqituvchi {i:02d}")
        ws = self.export()["2,3"]
        self.assertEqual(ws["B17"].value, "Oqituvchi 15")
        self.assertEqual(ws["C20"].value, "=SUM(C3:C19)")


class PatentTypesTest(ExportTestBase):
    def test_certificates_go_only_to_license_column(self):
        classic = self.teacher("Aliyev Ali")
        self.patent(classic)
        self.patent(classic, Patent.PatentType.SOFTWARE_CERTIFICATE, license=True)
        self.patent(classic, Patent.PatentType.DATABASE_CERTIFICATE)          # без договора — не в счёт
        only_cert = self.teacher("Boboyev Bobur")
        self.patent(only_cert, Patent.PatentType.SOFTWARE_CERTIFICATE)          # нигде не учитывается
        licensed_cert = self.teacher("Valiyev Vali")
        self.patent(licensed_cert, Patent.PatentType.DATABASE_CERTIFICATE, license=True)

        ws = self.export()["2,4"]
        rows = [tuple(ws.cell(row=r, column=c).value for c in (2, 3, 4)) for r in (3, 4, 5)]
        self.assertEqual(rows, [("Aliyev Ali", 1, 1), ("Valiyev Vali", 0, 1), (None, None, None)])


class CoAuthorShareTest(ExportTestBase):
    def test_same_article_entered_twice_is_split(self):
        # ikki kafedradosh bitta maqolani N=1 bilan kiritgan — 2,1 da har biriga 0.5
        for name in ("Aliyev Ali", "Valiyev Vali"):
            Publication.objects.create(teacher=self.teacher(name), reporting_period=self.period,
                                       article_title="Umumiy maqola", quartile="Q1")
        ws = self.export()["2,1"]
        self.assertEqual([(ws.cell(row=r, column=2).value, ws.cell(row=r, column=5).value) for r in (3, 4)],
                         [("Aliyev Ali", 0.5), ("Valiyev Vali", 0.5)])


class ArchiveTest(ExportTestBase):
    def test_documents_are_grouped_by_teacher_and_sheet(self):
        ali = self.teacher("Aliyev Ali", degree=self.phd)
        ali.diploma_file.save("diplom.pdf", ContentFile(b"%PDF diplom"))
        pub = Publication.objects.create(teacher=ali, reporting_period=self.period, article_title="Maqola")
        pub.article_file.save("maqola.pdf", ContentFile(b"%PDF maqola"))
        patent = self.patent(ali, Patent.PatentType.SOFTWARE_CERTIFICATE)
        patent.document_file.save("guvohnoma.pdf", ContentFile(b"%PDF guvohnoma"))
        funding = Funding.objects.create(teacher=ali, reporting_period=self.period,
                                         source_type=Funding.SourceType.STATE_PROGRAM, amount=1000)
        funding.document_file.save("shartnoma.pdf", ContentFile(b"%PDF shartnoma"))

        bobur = self.teacher("Boboyev Bobur", degree=self.phd)            # диплом не загружен
        Funding.objects.create(teacher=bobur, reporting_period=self.period,
                               source_type=Funding.SourceType.STATE_PROGRAM, amount=500)

        archive, filename = build_package(self.dep, self.period)
        self.assertEqual(filename, "Arxiv_TK_2026-mart.zip")
        with zipfile.ZipFile(archive) as zf:
            names = set(zf.namelist())
            register = openpyxl.load_workbook(io.BytesIO(zf.read("Arxiv_TK_2026-mart/Hujjatlar_royxati.xlsx")))

        root = "Arxiv_TK_2026-mart"
        self.assertIn(f"{root}/Reyting_TK_2026-mart.xlsx", names)
        for folder, file in [("Diplomlar", "diplom"), ("Maqolalar", "maqola"),
                             ("Patentlar va guvohnomalar", "guvohnoma"),
                             ("Mablag'lar", "shartnoma")]:
            self.assertTrue(any(n.startswith(f"{root}/Aliyev Ali/{folder}/{file}") for n in names), folder)
        # у Бобура ни одного файла — ни одной его папки в архиве
        self.assertFalse(any("Boboyev Bobur" in n for n in names))

        rows = list(register["Hujjatlar"].iter_rows(min_row=2, values_only=True))
        bobur_rows = {(r[2], r[5]) for r in rows if r[1] == "Boboyev Bobur"}
        self.assertEqual(bobur_rows, {("1,1", "fayl yuklanmagan"), ("2,6", "fayl yuklanmagan")})
        self.assertEqual(len([r for r in rows if r[1] == "Aliyev Ali"]), 4)
