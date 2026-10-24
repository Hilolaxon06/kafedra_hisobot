"""
Расширение секций листа формы под реальное число строк.

В шаблоне под каждую секцию заготовлено фиксированное число строк (14 на
основной штат, 1 на внутренних совместителей, 3 на патенты и т.п.). Если
людей больше — строки нужно вставить, иначе лишние молча пропадут из
официального отчёта.

Почему не просто ws.insert_rows(): openpyxl сдвигает только сами ячейки, а
объединённые ячейки (маркер «Ichki o'rindosh», примечание «*… ilova
qilinadi», подпись «Rektor … (imzo)»), высоты строк и ссылки в формулах
(«Jami: =SUM(F3:F16)») остаются на старых местах и ломают лист. Здесь всё это
сдвигается вместе со строками, а новые строки получают оформление последней
строки секции.
"""

import re
from copy import copy

from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter

# A1, $A$1, A1:B5 — ссылки внутри формулы
_REF_RE = re.compile(
    r"(?<![A-Za-z_\d])(\$?[A-Z]{1,3}\$?)(\d+)(?::(\$?[A-Z]{1,3}\$?)(\d+))?(?![\d(])"
)


def _shift_formula(formula, after_row, count):
    """Сдвигает ссылки ниже after_row на count строк; диапазон, который
    заканчивается на after_row или ниже, а начинается не ниже, — растягивает
    (новые строки попадают в «Jami» своей секции)."""

    def repl(m):
        col1, row1, col2, row2 = m.group(1), int(m.group(2)), m.group(3), m.group(4)
        if row1 > after_row:
            row1 += count
        if col2 is None:
            return f"{col1}{row1}"
        row2 = int(row2)
        if row2 >= after_row:
            row2 += count
        return f"{col1}{row1}:{col2}{row2}"

    return _REF_RE.sub(repl, formula)


def insert_rows_after(ws, row, count):
    """Вставляет count строк сразу под row, копируя оформление строки row."""
    if count <= 0:
        return

    # объединения ниже точки вставки — снять и потом поставить со сдвигом;
    # объединения внутри строки-прототипа — повторить на новых строках
    below, proto_merges = [], []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row > row:
            below.append((rng.min_row, rng.min_col, rng.max_row, rng.max_col))
            ws.unmerge_cells(str(rng))
        elif rng.min_row == rng.max_row == row:
            proto_merges.append((rng.min_col, rng.max_col))

    heights = {r: dim.height for r, dim in ws.row_dimensions.items() if r > row}
    for r in heights:
        ws.row_dimensions[r].height = None

    ws.insert_rows(row + 1, count)

    for r, height in heights.items():
        ws.row_dimensions[r + count].height = height
    for min_row, min_col, max_row, max_col in below:
        ws.merge_cells(start_row=min_row + count, start_column=min_col,
                       end_row=max_row + count, end_column=max_col)

    proto_height = ws.row_dimensions[row].height
    for new_row in range(row + 1, row + count + 1):
        for col in range(1, ws.max_column + 1):
            source = ws.cell(row=row, column=col)
            if source.has_style and not isinstance(source, MergedCell):
                ws.cell(row=new_row, column=col)._style = copy(source._style)
        ws.row_dimensions[new_row].height = proto_height
        for min_col, max_col in proto_merges:
            ws.merge_cells(f"{get_column_letter(min_col)}{new_row}:{get_column_letter(max_col)}{new_row}")

    for cells in ws.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = _shift_formula(cell.value, row, count)


class SectionRows:
    """Секции одного листа, перечисляемые сверху вниз. Номера строк задаются
    как в шаблоне; если секция выросла, все следующие секции сдвигаются
    автоматически."""

    def __init__(self, ws):
        self.ws = ws
        self.shift = 0

    def section(self, start, end, needed):
        """Возвращает фактические (start, end) секции, в которой не меньше
        needed строк."""
        start, end = start + self.shift, end + self.shift
        extra = needed - (end - start + 1)
        if extra > 0:
            insert_rows_after(self.ws, end, extra)
            end += extra
            self.shift += extra
        return start, end
