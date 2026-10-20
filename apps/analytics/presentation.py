"""
Yillik hisobot bo'yicha PowerPoint (.pptx) taqdimot.

Diagrammalar — PowerPoint'ning o'z (tahrirlanadigan) diagrammalari, rasm
emas: taqdimotni ochgandan keyin raqam yoki rangni o'zgartirish mumkin.
Slaydlar: titul → asosiy ko'rsatkichlar → reyting → pie-diagrammalar →
yillar dinamikasi → ball tarkibi → xulosa.
"""

from io import BytesIO

from django.utils import timezone
from django.utils.text import slugify
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from apps.core.models import Department

from . import stats

NAVY = RGBColor(0x14, 0x1B, 0x2D)
BLUE = RGBColor(0x2F, 0x5B, 0xEA)
INK = RGBColor(0x1B, 0x22, 0x33)
MUTED = RGBColor(0x6B, 0x74, 0x89)
SOFT = RGBColor(0xEA, 0xF0, 0xFF)
LINE = RGBColor(0xE6, 0xE9, 0xF0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SERIES = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300"]

W, H = Inches(13.333), Inches(7.5)


def _num(value):
    if isinstance(value, float) and not value.is_integer():
        return f"{value:,.2f}".rstrip("0").rstrip(".").replace(",", " ")
    return f"{int(value):,}".replace(",", " ")


def _text(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def _rect(slide, x, y, w, h, fill, line=None):
    shape = slide.shapes.add_shape(1, x, y, w, h)  # 1 = rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    shape.shadow.inherit = False
    return shape


def _slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # bo'sh maket
    _rect(slide, 0, 0, W, Inches(0.12), BLUE)
    _text(slide, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), title, size=28, bold=True)
    if subtitle:
        _text(slide, Inches(0.6), Inches(0.95), Inches(12), Inches(0.4), subtitle, size=14, color=MUTED)
    return slide


def _footer(slide, label, number):
    _text(slide, Inches(0.6), Inches(7.0), Inches(9), Inches(0.3), label, size=10, color=MUTED)
    _text(slide, Inches(11.7), Inches(7.0), Inches(1), Inches(0.3), str(number), size=10, color=MUTED,
          align=PP_ALIGN.RIGHT)


def _table(slide, x, y, w, header, rows, col_widths=None, highlight=None, size=12):
    shape = slide.shapes.add_table(len(rows) + 1, len(header), x, y, w, Inches(0.4) * (len(rows) + 1))
    table = shape.table
    if col_widths:
        total = sum(col_widths)
        for i, cw in enumerate(col_widths):
            table.columns[i].width = Emu(int(w * cw / total))
    for r, values in enumerate([header] + rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = value if isinstance(value, str) else _num(value)
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if c == 0 or (c == 1 and isinstance(value, str)) else PP_ALIGN.RIGHT
            font = para.runs[0].font if para.runs else para.font
            font.size = Pt(size)
            font.bold = r == 0 or (highlight is not None and r - 1 in highlight)
            font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if r == 0 else (SOFT if highlight and r - 1 in highlight else
                                                             (WHITE if r % 2 else RGBColor(0xF4, 0xF6, 0xFA)))
    return table


def _pie(slide, x, y, w, h, title, slices, unit=""):
    _text(slide, x, y, w, Inches(0.45), title, size=16, bold=True)
    if not slices:
        _rect(slide, x, y + Inches(0.55), w, h - Inches(0.55), RGBColor(0xF4, 0xF6, 0xFA), LINE)
        _text(slide, x, y + h / 2, w, Inches(0.4), "Ma’lumot yo‘q", size=14, color=MUTED, align=PP_ALIGN.CENTER)
        return
    data = CategoryChartData()
    data.categories = [name for name, _ in slices]
    data.add_series(title, [value for _, value in slices])
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y + Inches(0.45), w, h - Inches(0.45), data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.RIGHT
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(12)
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.show_percentage = True
    labels.show_value = False
    labels.number_format = "0%"
    labels.number_format_is_linked = False
    labels.font.size = Pt(11)
    labels.font.color.rgb = WHITE
    for i, point in enumerate(plot.series[0].points):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(SERIES[i % len(SERIES)])
        point.format.line.color.rgb = WHITE


def _kpis(total):
    return [
        ("O‘qituvchilar", total["staff"], ""),
        ("Ilmiy darajalilar ulushi", total["degree_share"], "%"),
        ("Maqolalar (ulush)", total["publications"], ""),
        ("Scopus / WoS", total["scopus_wos"], ""),
        ("H-indeks (jami)", total["hindex"], ""),
        ("Patentlar", total["patents"], ""),
        ("Mablag‘lar, mln so‘m", round(total["funding"] / 1_000_000, 1), ""),
        ("Himoyalar (PhD/DSc)", total["defenses"], ""),
    ]


def _conclusions(total, series, year, rows):
    lines = []
    prev = next((t for y, t in series if y == year - 1), None)
    if prev:
        for key, label in [("publications", "maqolalar"), ("scopus_wos", "Scopus/WoS maqolalari"),
                           ("patents", "patentlar"), ("hindex", "jami h-indeks")]:
            before, now = prev[key], total[key]
            if before or now:
                diff = now - before
                trend = "o‘sdi" if diff > 0 else "kamaydi" if diff < 0 else "o‘zgarmadi"
                lines.append(f"{year - 1}-yilga nisbatan {label} {trend}: {_num(before)} → {_num(now)}")
    else:
        lines.append(f"{year - 1}-yil uchun ma’lumot yo‘q — dinamika keyingi yildan ko‘rinadi.")
    lines.append(f"Ilmiy darajaga ega o‘qituvchilar ulushi: {_num(total['degree_share'])}%")
    if rows:
        best = rows[0]
        lines.append(f"Eng yuqori reyting: {best['department'].name} — "
                     f"1 o‘qituvchiga {_num(best['score_per_teacher'])} ball")
    missing = [label for key, label in [("patents", "patentlar"), ("funding", "jalb qilingan mablag‘lar"),
                                        ("mobility", "xalqaro mobillik"), ("defenses", "himoyalar")]
               if not total[key]]
    if missing:
        lines.append("Kuchaytirish kerak bo‘lgan yo‘nalishlar: " + ", ".join(missing))
    return lines


def build_presentation(departments, year, user=None):
    rows, total = stats.overview(departments, year)
    series = stats.yearly_series(departments)
    all_rows = stats.ranking(Department.objects.order_by("name"), year)
    chosen = {d.pk for d in departments}
    # faqat shu yili hisobot davri bor kafedralar
    shown = [r["department"] for r in rows] or list(departments)
    full_names = ", ".join(d.name for d in shown)
    names = shown[0].name if len(shown) == 1 else f"{len(shown)} ta kafedra"

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    footer = f"{names} · {year}-yil hisoboti"
    n = 1

    # 1. Titul
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, W, H, NAVY)
    _rect(slide, Inches(0.8), Inches(2.3), Inches(0.12), Inches(2.2), BLUE)
    _text(slide, Inches(1.2), Inches(2.2), Inches(11), Inches(1), f"{year}-yil yillik hisoboti",
          size=44, bold=True, color=WHITE)
    _text(slide, Inches(1.2), Inches(3.3), Inches(11), Inches(1.2), full_names, size=24 if len(departments) == 1 else 18, color=RGBColor(0xC9, 0xD1, 0xE4))
    author = ""
    if user is not None:
        teacher = getattr(user, "teacher_profile", None)
        author = teacher.full_name if teacher else (user.get_full_name() or user.username)
    _text(slide, Inches(1.2), Inches(6.2), Inches(11), Inches(0.5),
          " · ".join(filter(None, [author, timezone.localdate().strftime("%d.%m.%Y")])),
          size=14, color=RGBColor(0xAA, 0xB4, 0xCC))

    # 2. Asosiy ko'rsatkichlar
    n += 1
    slide = _slide(prs, "Asosiy ko‘rsatkichlar", f"{year}-yil, {names}")
    tile_w, tile_h, gap = Inches(2.85), Inches(2.2), Inches(0.25)
    for i, (label, value, suffix) in enumerate(_kpis(total)):
        x = Inches(0.6) + (tile_w + gap) * (i % 4)
        y = Inches(1.7) + (tile_h + gap) * (i // 4)
        _rect(slide, x, y, tile_w, tile_h, WHITE, LINE)
        _rect(slide, x, y, Inches(0.08), tile_h, BLUE)
        _text(slide, x + Inches(0.3), y + Inches(0.35), tile_w - Inches(0.4), Inches(0.9),
              _num(value) + suffix, size=36, bold=True)
        _text(slide, x + Inches(0.3), y + Inches(1.35), tile_w - Inches(0.4), Inches(0.6), label, size=14, color=MUTED)
    _footer(slide, footer, n)

    # 3. Reyting
    if all_rows:
        n += 1
        slide = _slide(prs, "Kafedralar reytingi", f"{year}-yil · 1 o‘qituvchiga ball bo‘yicha")
        shown = all_rows[:12]
        _table(slide, Inches(0.6), Inches(1.6), Inches(12.1),
               ["O‘rin", "Kafedra", "O‘qit.", "Maqolalar", "Scopus/WoS", "Patent", "Daraja %", "Jami ball", "1 o‘qit.ga"],
               [[str(r["place"]), r["department"].name, r["staff"], r["publications"], r["scopus_wos"],
                 r["patents"], r["degree_share"], r["score"], r["score_per_teacher"]] for r in shown],
               col_widths=[0.7, 4, 0.9, 1.1, 1.2, 0.9, 1, 1.1, 1.1],
               highlight={i for i, r in enumerate(shown) if r["department"].pk in chosen}, size=12)
        _footer(slide, footer, n)

    # 4. Pie-diagrammalar — har slaydda ikkitadan
    charts = [c for c in stats.charts(total) if c[1]]  # bo'sh diagrammalar taqdimotga kirmaydi
    for i in range(0, len(charts), 2):
        n += 1
        slide = _slide(prs, "Statistika", f"{year}-yil · ulushlar")
        for j, (title, slices, unit) in enumerate(charts[i:i + 2]):
            _pie(slide, Inches(0.6) + Inches(6.2) * j, Inches(1.6), Inches(5.9), Inches(5.1), title, slices, unit)
        _footer(slide, footer, n)

    # 5. Yillar dinamikasi
    if len(series) > 1:
        n += 1
        slide = _slide(prs, "Yillar dinamikasi", "Asosiy ko‘rsatkichlarning yillar bo‘yicha o‘zgarishi")
        data = CategoryChartData()
        data.categories = [str(y) for y, _ in series]
        data.add_series("Maqolalar", [t["publications"] for _, t in series])
        data.add_series("Scopus / WoS", [t["scopus_wos"] for _, t in series])
        data.add_series("Patentlar", [t["patents"] for _, t in series])
        chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.5),
                                       Inches(7.2), Inches(5.3), data).chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.plots[0].has_data_labels = True
        chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        chart.plots[0].data_labels.font.size = Pt(10)
        for k, s in enumerate(chart.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = RGBColor.from_string(SERIES[k])
        _table(slide, Inches(8.1), Inches(1.6), Inches(4.6), ["Ko‘rsatkich"] + [str(y) for y, _ in series[-3:]],
               [[label] + [t[key] for _, t in series[-3:]] for key, label in [
                   ("staff", "O‘qituvchilar"), ("degree_share", "Daraja, %"), ("publications", "Maqolalar"),
                   ("scopus_wos", "Scopus/WoS"), ("hindex", "H-indeks"), ("patents", "Patentlar"),
                   ("defenses", "Himoyalar"), ("mobility", "Mobillik")]], size=11)
        _footer(slide, footer, n)

    # 6. Ball tarkibi
    if rows:
        n += 1
        slide = _slide(prs, "Reyting balli tarkibi", "Har bir yo‘nalish bo‘yicha to‘plangan ball")
        parts = list(rows[0]["score_parts"].keys())
        _table(slide, Inches(0.6), Inches(1.6), Inches(12.1), ["Kafedra"] + parts + ["Jami"],
               [[r["department"].name] + list(r["score_parts"].values()) + [r["score"]] for r in rows[:12]],
               col_widths=[4] + [1.2] * len(parts) + [1.1], size=12)
        _footer(slide, footer, n)

    # 7. Xulosa
    n += 1
    slide = _slide(prs, "Xulosa", f"{year}-yil yakunlari")
    _text(slide, Inches(0.8), Inches(1.7), Inches(11.8), Inches(5),
          ["•  " + line for line in _conclusions(total, series, year, all_rows)], size=18)
    _footer(slide, footer, n)

    buffer = BytesIO()
    prs.save(buffer)
    label = (slugify(departments[0].name) if len(departments) == 1 else "") or "kafedralar"
    return buffer.getvalue(), f"taqdimot_{label}_{year}.pptx"
